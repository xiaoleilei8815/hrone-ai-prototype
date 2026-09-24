#!/usr/bin/env python3
"""Compare cached results for every formula cell in two workbooks."""

import json
import math
import sys
from itertools import zip_longest

from openpyxl import load_workbook


def equal(left, right, number_format):
    if left is None or right is None:
        return left is right
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        absolute_tolerance = 5e-8 if "%" in str(number_format or "") else 1e-5
        return math.isclose(
            float(left), float(right), rel_tol=1e-9, abs_tol=absolute_tolerance
        )
    return left == right


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: compare-formula-cache.py <source.xlsx> <recalculated.xlsx>")

    source_path, recalculated_path = sys.argv[1:]
    formulas = load_workbook(source_path, data_only=False, read_only=True)
    source_values = load_workbook(source_path, data_only=True, read_only=True)
    recalculated_values = load_workbook(recalculated_path, data_only=True, read_only=True)
    differences = []
    difference_count = 0
    formula_count = 0
    max_absolute_difference = 0.0
    max_percentage_difference = 0.0

    for sheet_name in formulas.sheetnames:
        formula_sheet = formulas[sheet_name]
        source_sheet = source_values[sheet_name]
        recalculated_sheet = recalculated_values[sheet_name]
        rows = zip_longest(
            formula_sheet.iter_rows(),
            source_sheet.iter_rows(),
            recalculated_sheet.iter_rows(),
            fillvalue=(),
        )
        for formula_row, source_row, recalculated_row in rows:
            for index, cell in enumerate(formula_row):
                if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                    continue
                formula_count += 1
                left = source_row[index].value if index < len(source_row) else None
                right = recalculated_row[index].value if index < len(recalculated_row) else None
                if isinstance(left, (int, float)) and isinstance(right, (int, float)):
                    absolute_difference = abs(float(left) - float(right))
                    max_absolute_difference = max(max_absolute_difference, absolute_difference)
                    if "%" in str(cell.number_format or ""):
                        max_percentage_difference = max(
                            max_percentage_difference, absolute_difference
                        )
                if not equal(left, right, cell.number_format):
                    difference_count += 1
                    if len(differences) < 20:
                        differences.append(
                            {
                                "sheet": sheet_name,
                                "cell": cell.coordinate,
                                "source": left,
                                "recalculated": right,
                            }
                        )

    result = {
        "status": "PASS" if difference_count == 0 else "FAIL",
        "formula_count": formula_count,
        "difference_count": difference_count,
        "max_absolute_difference": max_absolute_difference,
        "max_percentage_difference": max_percentage_difference,
        "difference_examples": differences,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
    if difference_count:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
