from __future__ import annotations

import json
import math
import sys
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parents[4]
REFERENCE = (
    ROOT
    / "outputs"
    / "europe_supplier_performance_q2_2026_final"
    / "欧洲供应商履约率报表_2026年4-6月_最终严谨版.xlsx"
)
EXPECTED_SHEETS = [
    "管理摘要",
    "主指标_原始兑现率",
    "对照_超发校正率",
    "供应商月度明细",
    "仓库满足与发单质量",
    "供应商责任底表",
    "仓库单元底表",
    "三表对账",
    "口径与质量审查",
    "源文件清单",
]
CORE_SHEETS = [
    "主指标_原始兑现率",
    "对照_超发校正率",
    "供应商月度明细",
    "仓库满足与发单质量",
    "供应商责任底表",
    "仓库单元底表",
]
EXPECTED_TABLES = {
    "主指标_原始兑现率": ("RawFulfillmentScorecard", "A4:L155"),
    "对照_超发校正率": ("AdjustedFulfillmentScorecard", "A4:M155"),
    "供应商月度明细": ("SupplierMonthlyMetrics", "A4:O155"),
    "仓库满足与发单质量": ("WarehouseAndAllocationSummary", "A4:Q79"),
    "供应商责任底表": ("SupplierAtomicResponsibility", "A4:V1541"),
    "仓库单元底表": ("WarehouseAtomicFacts", "A4:S1321"),
    "三表对账": ("MonthlyThreeWayReconciliation", "A4:N21"),
    "源文件清单": ("SourceFileRegistry", "A4:F11"),
}
EPS = 1e-6


def equal_value(left: Any, right: Any) -> bool:
    if left in (None, "") and right in (None, ""):
        return True
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return math.isclose(float(left), float(right), abs_tol=EPS, rel_tol=EPS)
    return left == right


def color_signature(color: Any) -> tuple[Any, ...] | None:
    if color is None:
        return None
    return (color.type, color.rgb, color.indexed, color.theme, color.tint)


def style_signature(cell: Any) -> tuple[Any, ...]:
    border_style = lambda side: side.style if side is not None else None
    return (
        cell.font.name,
        cell.font.sz,
        cell.font.bold,
        cell.font.italic,
        color_signature(cell.font.color),
        cell.fill.fill_type,
        color_signature(cell.fill.fgColor),
        cell.number_format,
        cell.alignment.horizontal,
        cell.alignment.vertical,
        cell.alignment.wrap_text,
        border_style(cell.border.left),
        border_style(cell.border.right),
        border_style(cell.border.top),
        border_style(cell.border.bottom),
    )


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("usage: verify-export.py <export.xlsx>")
    output = Path(sys.argv[1])
    reference_values = openpyxl.load_workbook(REFERENCE, data_only=True)
    output_values = openpyxl.load_workbook(output, data_only=True)
    output_formulas = openpyxl.load_workbook(output, data_only=False)

    differences: list[dict[str, Any]] = []
    if output_values.sheetnames != EXPECTED_SHEETS:
        differences.append(
            {
                "type": "sheet_names",
                "expected": EXPECTED_SHEETS,
                "actual": output_values.sheetnames,
            }
        )

    comparison_count = 0
    for sheet_name in CORE_SHEETS:
        expected = reference_values[sheet_name]
        actual = output_values[sheet_name]
        max_row = max(expected.max_row, actual.max_row)
        max_column = max(expected.max_column, actual.max_column)
        for row in range(1, max_row + 1):
            for column in range(1, max_column + 1):
                comparison_count += 1
                left = expected.cell(row, column).value
                right = actual.cell(row, column).value
                if not equal_value(left, right):
                    differences.append(
                        {
                            "type": "core_value",
                            "sheet": sheet_name,
                            "cell": openpyxl.utils.get_column_letter(column)
                            + str(row),
                            "expected": left,
                            "actual": right,
                        }
                    )

    dashboard_cells = [
        "B6",
        "C6",
        "D6",
        "E6",
        "B7",
        "C7",
        "D7",
        "E7",
        "B8",
        "C8",
        "D8",
        "E8",
        "B9",
        "C9",
        "D9",
        "E9",
        "B10",
        "C10",
        "D10",
        "E10",
        "B13",
        "D13",
        "F13",
        "H13",
        "B14",
        "D14",
        "F14",
        "H14",
    ]
    for cell in dashboard_cells:
        comparison_count += 1
        left = reference_values["管理摘要"][cell].value
        right = output_values["管理摘要"][cell].value
        if not equal_value(left, right):
            differences.append(
                {
                    "type": "dashboard_value",
                    "sheet": "管理摘要",
                    "cell": cell,
                    "expected": left,
                    "actual": right,
                }
            )

    formula_count = 0
    formula_errors: list[dict[str, Any]] = []
    for sheet in output_formulas.worksheets:
        cached_sheet = output_values[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
                    cached = cached_sheet[cell.coordinate].value
                    if isinstance(cached, str) and cached.startswith("#"):
                        formula_errors.append(
                            {
                                "sheet": sheet.title,
                                "cell": cell.coordinate,
                                "value": cached,
                            }
                        )

    table_results = {}
    for sheet_name, (table_name, expected_ref) in EXPECTED_TABLES.items():
        sheet = output_formulas[sheet_name]
        table = sheet.tables.get(table_name)
        actual_ref = table.ref if table else None
        table_results[table_name] = actual_ref
        if actual_ref != expected_ref:
            differences.append(
                {
                    "type": "table_ref",
                    "sheet": sheet_name,
                    "table": table_name,
                    "expected": expected_ref,
                    "actual": actual_ref,
                }
            )

    if formula_count != 26196:
        differences.append(
            {
                "type": "formula_count",
                "expected": 26196,
                "actual": formula_count,
            }
        )
    if formula_errors:
        differences.extend({"type": "formula_error", **item} for item in formula_errors)
    if len(output_formulas["管理摘要"]._charts) != 1:
        differences.append(
            {
                "type": "chart_count",
                "expected": 1,
                "actual": len(output_formulas["管理摘要"]._charts),
            }
        )

    style_checks = [
        ("管理摘要", "A1"),
        ("管理摘要", "A5"),
        ("主指标_原始兑现率", "A4"),
        ("主指标_原始兑现率", "F5"),
        ("主指标_原始兑现率", "A6"),
        ("对照_超发校正率", "M5"),
        ("供应商月度明细", "I5"),
        ("仓库满足与发单质量", "N5"),
        ("供应商责任底表", "R5"),
        ("仓库单元底表", "O5"),
    ]
    for sheet_name, cell in style_checks:
        expected_style = style_signature(reference_values[sheet_name][cell])
        actual_style = style_signature(output_values[sheet_name][cell])
        if expected_style != actual_style:
            differences.append(
                {
                    "type": "style",
                    "sheet": sheet_name,
                    "cell": cell,
                    "expected": expected_style,
                    "actual": actual_style,
                }
            )

    report = {
        "status": "PASS" if not differences else "FAIL",
        "workbook": str(output),
        "sheet_count": len(output_values.sheetnames),
        "formula_count": formula_count,
        "formula_error_count": len(formula_errors),
        "value_comparison_count": comparison_count,
        "difference_count": len(differences),
        "difference_examples": differences[:30],
        "tables": table_results,
        "chart_count": len(output_formulas["管理摘要"]._charts),
    }
    print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
    reference_values.close()
    output_values.close()
    output_formulas.close()
    if report["status"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
