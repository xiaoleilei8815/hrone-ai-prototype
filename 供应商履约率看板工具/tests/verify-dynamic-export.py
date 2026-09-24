from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parents[4]
REFERENCE = (
    ROOT
    / "outputs"
    / "europe_supplier_performance_q2_2026_final"
    / "欧洲供应商履约率报表_2026年4-6月_最终严谨版.xlsx"
)
EXPECTED_SHEETS = [
    "管理摘要",
    "主指标_原始兑现率",
    "对照_超发校正率",
    "供应商月度明细",
    "仓库满足与发单质量",
    "供应商责任底表",
    "仓库单元底表",
    "三表对账",
    "口径与质量审查",
    "源文件清单",
]
EXPECTED_TABLES = {
    "主指标_原始兑现率": ("RawFulfillmentScorecard", None),
    "对照_超发校正率": ("AdjustedFulfillmentScorecard", None),
    "供应商月度明细": ("SupplierMonthlyMetrics", "O"),
    "仓库满足与发单质量": ("WarehouseAndAllocationSummary", "Q"),
    "供应商责任底表": ("SupplierAtomicResponsibility", "V"),
    "仓库单元底表": ("WarehouseAtomicFacts", "S"),
    "三表对账": ("MonthlyThreeWayReconciliation", "N"),
    "源文件清单": ("SourceFileRegistry", "F"),
}


def color_signature(color: Any) -> tuple[Any, ...] | None:
    if color is None:
        return None
    return (color.type, color.rgb, color.indexed, color.theme, color.tint)


def style_signature(cell: Any) -> tuple[Any, ...]:
    border_style = lambda side: side.style if side is not None else None
    return (
        cell.font.name,
        cell.font.sz,
        cell.font.bold,
        cell.font.italic,
        color_signature(cell.font.color),
        cell.fill.fill_type,
        color_signature(cell.fill.fgColor),
        cell.number_format,
        cell.alignment.horizontal,
        cell.alignment.vertical,
        cell.alignment.wrap_text,
        border_style(cell.border.left),
        border_style(cell.border.right),
        border_style(cell.border.top),
        border_style(cell.border.bottom),
    )


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("usage: verify-dynamic-export.py <export.xlsx>")
    output = Path(sys.argv[1])
    reference = openpyxl.load_workbook(REFERENCE, data_only=True)
    values = openpyxl.load_workbook(output, data_only=True)
    formulas = openpyxl.load_workbook(output, data_only=False)
    errors: list[dict[str, Any]] = []

    if formulas.sheetnames != EXPECTED_SHEETS:
        errors.append({"type": "sheet_names", "actual": formulas.sheetnames})

    formula_count = 0
    formula_errors = 0
    precedence_hazards = 0
    for sheet in formulas.worksheets:
        cached = values[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
                    compact_formula = cell.value.replace(" ", "")
                    if (
                        "+SUMIFS" in compact_formula
                        and "/" in compact_formula
                        and ")/(" not in compact_formula
                    ):
                        precedence_hazards += 1
                        errors.append(
                            {
                                "type": "formula_precedence",
                                "sheet": sheet.title,
                                "cell": cell.coordinate,
                            }
                        )
                    if "#REF!" in cell.value:
                        errors.append(
                            {"type": "formula_ref", "sheet": sheet.title, "cell": cell.coordinate}
                        )
                    cached_value = cached[cell.coordinate].value
                    if isinstance(cached_value, str) and cached_value.startswith("#"):
                        formula_errors += 1
                        errors.append(
                            {
                                "type": "formula_error",
                                "sheet": sheet.title,
                                "cell": cell.coordinate,
                                "value": cached_value,
                            }
                        )

    table_results: dict[str, str | None] = {}
    for sheet_name, (table_name, configured_end_column) in EXPECTED_TABLES.items():
        sheet = formulas[sheet_name]
        table = sheet.tables.get(table_name)
        table_results[table_name] = table.ref if table else None
        if not table:
            errors.append({"type": "missing_table", "table": table_name})
            continue
        end_column = configured_end_column
        if end_column is None:
            last_header_column = max(
                cell.column for cell in sheet[4] if cell.value not in (None, "")
            )
            end_column = openpyxl.utils.get_column_letter(last_header_column)
        if not table.ref.startswith("A4:") or not table.ref.split(":")[1].startswith(end_column):
            errors.append({"type": "table_ref", "table": table_name, "actual": table.ref})

    style_checks = [
        ("管理摘要", "A1"),
        ("管理摘要", "A5"),
        ("主指标_原始兑现率", "A4"),
        ("主指标_原始兑现率", "F5"),
        ("对照_超发校正率", "F5"),
        ("供应商月度明细", "I5"),
        ("仓库满足与发单质量", "N5"),
        ("供应商责任底表", "R5"),
        ("仓库单元底表", "O5"),
    ]
    for sheet_name, coordinate in style_checks:
        if style_signature(reference[sheet_name][coordinate]) != style_signature(
            values[sheet_name][coordinate]
        ):
            errors.append(
                {"type": "style", "sheet": sheet_name, "cell": coordinate}
            )

    chart_count = len(formulas["管理摘要"]._charts)
    if chart_count != 1:
        errors.append({"type": "chart_count", "actual": chart_count})

    report = {
        "status": "PASS" if not errors else "FAIL",
        "workbook": str(output),
        "sheet_count": len(formulas.sheetnames),
        "formula_count": formula_count,
        "formula_error_count": formula_errors,
        "formula_precedence_hazard_count": precedence_hazards,
        "tables": table_results,
        "chart_count": chart_count,
        "error_count": len(errors),
        "error_examples": errors[:20],
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    reference.close()
    values.close()
    formulas.close()
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
